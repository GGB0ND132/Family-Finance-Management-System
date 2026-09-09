"""CSV/XLSX 文件解析器。

仅做“读取文件 -> 得到 [{列名: 值}, ...]”的机械转换，不包含业务校验；
业务校验在 `service.py` 中完成。CSV 兼容 UTF-8（含 BOM）与 GBK 编码。
"""

import csv
import io

from app.core.exceptions import BadRequestError


def parse_file(file_bytes: bytes, filename: str) -> list[dict]:
    """解析上传文件，返回行数据列表，每行形如 `{列名: 单元格文本}`。"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "csv":
        return _parse_csv(file_bytes)
    elif ext == "xlsx":
        return _parse_xlsx(file_bytes)
    else:
        raise BadRequestError(f"不支持的文件格式: .{ext}")


def _parse_csv(data: bytes) -> list[dict]:
    text = _decode_text(data)
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise BadRequestError("CSV 文件缺少表头")
    headers = [h.strip() for h in reader.fieldnames]
    rows: list[dict] = []
    for raw in reader:
        rows.append({header: (raw.get(header) or "").strip() for header in headers})
    return rows


def _parse_xlsx(data: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - 依赖缺失保护
        raise BadRequestError("服务端未安装 openpyxl，无法解析 XLSX") from exc

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers_row = next(rows_iter, None)
        if headers_row is None:
            raise BadRequestError("XLSX 文件缺少表头")
        headers = [str(h).strip() if h is not None else "" for h in headers_row]
        if not any(headers):
            raise BadRequestError("XLSX 文件缺少表头")

        rows: list[dict] = []
        for values in rows_iter:
            if values is None or all(v is None or str(v).strip() == "" for v in values):
                continue
            cell = {header: _cell_to_str(values[idx]) for idx, header in enumerate(headers)}
            rows.append(cell)
        return rows
    finally:
        wb.close()


def _decode_text(data: bytes) -> str:
    """优先按 UTF-8（含 BOM）解码，失败则回退 GBK。"""
    for encoding in ("utf-8-sig", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
